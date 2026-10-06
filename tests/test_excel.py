from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from services import excel


def _df():
    return pd.DataFrame({
        "Publicador": ["Ana", "=HYPERLINK(\"http://malo\")", "Nico"],
        "Grupo": ["Grupo 1", "Grupo 1", "Grupo 2"],
        "Mes": ["2026-10", "2026-10", "2026-10"],
        "Privilegios": ["Precursor Regular", "Publicador", "Inactivo"],
        "Informe": ["50 horas", "Sí participé", "INACTIVO (No predicó)"],
        "Cursos Bíblicos": [2, 0, 0],
        "Horas": [50, 0, 0],
    })


def test_preparar_exportacion_deja_blancos_los_inactivos_y_ceros():
    out = excel.preparar_exportacion(_df())
    assert list(out.columns) == excel.COLUMNAS_EXPORTACION
    assert pd.isna(out.loc[0, "Cursos bíblicos"]) is False and out.loc[0, "Cursos bíblicos"] == 2
    assert pd.isna(out.loc[1, "Cursos bíblicos"]) and pd.isna(out.loc[1, "Horas"])  # 0 -> celda vacía
    assert out.loc[2, "Mes"] == "" and out.loc[2, "Informe"] == ""                   # inactivo


def test_excel_tiene_tabla_filtrable_totales_y_resumen():
    contenido = excel.generar_excel(excel.preparar_exportacion(_df()), "Reporte general · 2026-10")
    libro = load_workbook(BytesIO(contenido))
    assert libro.sheetnames == ["Informes", "Resumen por grupo"]
    hoja = libro["Informes"]
    assert hoja["A1"].value == "Reporte general · 2026-10"
    assert list(hoja.tables) == ["TablaInformes"]
    assert hoja.tables["TablaInformes"].ref == "A4:G7"
    assert hoja.freeze_panes == "A5"
    # los totales usan SUBTOTAL: se recalculan al filtrar
    formulas = [c.value for c in hoja[3] if isinstance(c.value, str) and c.value.startswith("=SUBTOTAL")]
    assert len(formulas) == 2


def test_texto_con_igual_no_se_guarda_como_formula():
    hoja = load_workbook(BytesIO(excel.generar_excel(excel.preparar_exportacion(_df()), "t")))["Informes"]
    celda = hoja["A6"]
    assert celda.data_type == "s" and celda.value.startswith("=")


def test_excel_vacio_no_falla():
    vacio = excel.preparar_exportacion(_df().iloc[0:0])
    assert load_workbook(BytesIO(excel.generar_excel(vacio, "vacío")))["Informes"]["A1"].value == "vacío"
