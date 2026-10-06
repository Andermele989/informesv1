"""Exportación a Excel: tabla filtrable con totales que se recalculan al filtrar.

Diseño sobrio: título, tabla nativa de Excel (filtros, orden y filas alternas), fila de
totales con SUBTOTAL (respeta los filtros) y una hoja de resumen por grupo.
"""
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

_NAVY = "1E293B"
_FUENTE = "Calibri"
_BORDE_FINO = Side(style="thin", color="CBD5E1")

COLUMNAS_EXPORTACION = ["Nombre y apellido", "Privilegio", "Mes", "Informe", "Cursos bíblicos", "Grupo", "Horas"]


def preparar_exportacion(df: pd.DataFrame) -> pd.DataFrame:
    """Deja el DataFrame del dashboard con las columnas del Excel.

    Las filas de publicadores inactivos salen con Mes, Informe y Cursos en blanco.
    """
    out = pd.DataFrame({
        "Nombre y apellido": df["Publicador"],
        "Privilegio": df["Privilegios"],
        "Mes": df["Mes"],
        "Informe": df["Informe"],
        "Cursos bíblicos": df["Cursos Bíblicos"].where(df["Cursos Bíblicos"] > 0),  # 0 -> celda vacía
        "Grupo": df["Grupo"],
        "Horas": df["Horas"].where(df["Horas"] > 0),
    }).astype({"Mes": object, "Informe": object})
    inactivos = out["Privilegio"].str.lower() == "inactivo"
    out.loc[inactivos, ["Mes", "Informe"]] = ""
    out.loc[inactivos, ["Cursos bíblicos", "Horas"]] = None
    return out.reset_index(drop=True)


def _escribir(celda, valor) -> None:
    """Asigna un valor. Un texto que empiece por "=" se guarda como texto, nunca como fórmula."""
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        valor = None
    if hasattr(valor, "item"):  # numpy -> tipo nativo
        valor = valor.item()
    celda.value = valor
    if isinstance(valor, str) and valor.startswith("="):
        celda.data_type = "s"


def _ancho(encabezado: str, valores: pd.Series) -> float:
    mayor = max([len(str(v)) for v in valores.head(500) if v is not None and not pd.isna(v)] + [0])
    return max(11, min(max(len(encabezado) + 4, mayor + 2), 46))


def _hoja_tabla(ws, df: pd.DataFrame, titulo: str, nombre_tabla: str, totales: bool = True) -> None:
    ws.sheet_view.showGridLines = False
    n_cols, n_filas = len(df.columns), len(df)
    ultima = get_column_letter(n_cols)
    fila_totales, fila_enc, primera = 3, 4, 5
    ultima_fila = max(primera, primera + n_filas - 1)

    ws["A1"] = titulo
    ws["A1"].font = Font(name=_FUENTE, size=15, bold=True, color=_NAVY)
    ws["A2"] = f"Generado el {datetime.now():%d/%m/%Y %H:%M} · {n_filas} registros"
    ws["A2"].font = Font(name=_FUENTE, size=10, color="64748B")
    ws.row_dimensions[1].height = 26

    for j, columna in enumerate(df.columns, start=1):
        cab = ws.cell(row=fila_enc, column=j, value=str(columna))
        cab.font = Font(name=_FUENTE, size=11, bold=True, color="FFFFFF")
        cab.fill = PatternFill("solid", start_color=_NAVY)
        cab.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = _ancho(str(columna), df[columna])
    ws.row_dimensions[fila_enc].height = 24

    for i, fila in enumerate(df.itertuples(index=False), start=primera):
        for j, valor in enumerate(fila, start=1):
            celda = ws.cell(row=i, column=j)
            _escribir(celda, valor)
            numerico = pd.api.types.is_numeric_dtype(df.iloc[:, j - 1])
            celda.font = Font(name=_FUENTE, size=10)
            celda.alignment = Alignment(horizontal="center" if numerico else "left", vertical="center")
            celda.border = Border(bottom=_BORDE_FINO)
            if numerico:
                celda.number_format = "#,##0"

    # Totales: SUBTOTAL ignora las filas ocultas por el filtro.
    if totales and n_filas:
        ws.cell(row=fila_totales, column=1, value="Total (según filtro)").font = Font(name=_FUENTE, size=10, bold=True, color=_NAVY)
        for j, columna in enumerate(df.columns, start=1):
            if pd.api.types.is_numeric_dtype(df[columna]):
                letra = get_column_letter(j)
                celda = ws.cell(row=fila_totales, column=j, value=f"=SUBTOTAL(109,{letra}{primera}:{letra}{ultima_fila})")
                celda.font = Font(name=_FUENTE, size=10, bold=True, color=_NAVY)
                celda.alignment = Alignment(horizontal="center")
                celda.number_format = "#,##0"

    tabla = Table(displayName=nombre_tabla, ref=f"A{fila_enc}:{ultima}{ultima_fila}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tabla)
    ws.freeze_panes = f"A{primera}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{fila_enc}:{fila_enc}"


def _resumen_por_grupo(df: pd.DataFrame) -> pd.DataFrame:
    activos = df[df["Privilegio"].str.lower() != "inactivo"]
    if activos.empty:
        return pd.DataFrame(columns=["Grupo", "Informes", "Horas", "Cursos bíblicos", "Promedio de horas"])
    resumen = activos.groupby("Grupo").agg(
        Informes=("Nombre y apellido", "count"), Horas=("Horas", "sum"), Cursos=("Cursos bíblicos", "sum"),
    ).reset_index().rename(columns={"Cursos": "Cursos bíblicos"})
    resumen["Promedio de horas"] = (resumen["Horas"] / resumen["Informes"]).round(1)
    return resumen.astype({"Horas": int, "Cursos bíblicos": int})


def generar_excel(df: pd.DataFrame, titulo: str, hoja: str = "Informes") -> bytes:
    """Excel de informes (hoja filtrable + resumen por grupo)."""
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    _hoja_tabla(ws, df, titulo, "TablaInformes")
    if {"Grupo", "Horas"} <= set(df.columns) and not df.empty:
        resumen = wb.create_sheet("Resumen por grupo")
        _hoja_tabla(resumen, _resumen_por_grupo(df), f"Resumen por grupo · {titulo}", "TablaResumen")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar_excel_simple(df: pd.DataFrame, titulo: str, hoja: str) -> bytes:
    """Excel de una sola hoja filtrable (para el directorio de publicadores)."""
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    _hoja_tabla(ws, df, titulo, "TablaDatos")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
